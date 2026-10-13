#!/usr/bin/env python3
"""Render real Office output to PNG and a contact sheet. No synthetic PPT renderer fallback.
PowerPoint COM (Windows) or LibreOffice exports PDF; PDFium rasterizes it.
PDF inputs also supported. --probe reports renderer availability without launching Office.
"""
import argparse
import json
import math
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
import pypdfium2 as pdfium
from PIL import Image, ImageDraw
from _common import utf8_stdio


def find_soffice(explicit=None):
    candidates = [explicit, os.environ.get('SOFFICE'), shutil.which('soffice'), shutil.which('libreoffice'),
                  r'C:\Program Files\LibreOffice\program\soffice.exe',
                  r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
                  '/Applications/LibreOffice.app/Contents/MacOS/soffice']
    return next((str(Path(p).resolve()) for p in candidates if p and Path(p).is_file()), None)


def office_pdf(source, dest, engine='auto', soffice=None, timeout=120):
    errors = []
    if os.name == 'nt' and engine in ('auto', 'powerpoint'):
        # Paths travel as environment variables, never interpolate them into PowerShell code.
        script = r'''
$ErrorActionPreference = 'Stop'
$app = $null
$deck = $null
$hadPowerPoint = @(Get-Process POWERPNT -ErrorAction SilentlyContinue).Count -gt 0
$verifiedPowerPoint = $false
try {
  $server = $null
  foreach ($view in @([Microsoft.Win32.RegistryView]::Registry64, [Microsoft.Win32.RegistryView]::Registry32)) {
    $hive = [Microsoft.Win32.RegistryKey]::OpenBaseKey([Microsoft.Win32.RegistryHive]::ClassesRoot, $view)
    try {
      $prog = 'PowerPoint.Application'
      $key = $hive.OpenSubKey($prog + '\CurVer')
      if ($null -ne $key) { $version = $key.GetValue(''); $key.Close(); if ($version) { $prog = $version } }
      $key = $hive.OpenSubKey($prog + '\CLSID')
      if ($null -eq $key) { continue }
      $clsid = $key.GetValue(''); $key.Close()
      $key = $hive.OpenSubKey('CLSID\' + $clsid + '\LocalServer32')
      if ($null -eq $key) { continue }
      $candidate = [string]$key.GetValue(''); $key.Close()
      if ($candidate -match '(?i)powerpnt\.exe') { $server = $candidate; break }
    } finally { $hive.Close() }
  }
  if (-not $server) { throw 'No registered Microsoft PowerPoint executable in either registry view; not launching another application.' }
  $app = New-Object -ComObject PowerPoint.Application
  $appPath = [string]$app.Path
  if (-not (Test-Path -LiteralPath (Join-Path $appPath 'POWERPNT.EXE'))) {
    throw 'PowerPoint COM is redirected to another application; not a verified Microsoft renderer.'
  }
  $verifiedPowerPoint = $true
  $deck = $app.Presentations.Open($env:PBS_SOURCE, -1, 0, 0)
  $deck.SaveAs($env:PBS_DEST, 32)
} finally {
  if ($null -ne $deck) { $deck.Close(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($deck) }
  if ($null -ne $app) {
    if ($verifiedPowerPoint -and -not $hadPowerPoint) { $app.Quit() }
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($app)
  }
}
'''
        env = {**os.environ, 'PBS_SOURCE': str(source), 'PBS_DEST': str(dest)}
        try:
            proc = subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-Command', script], env=env,
                                  capture_output=True, timeout=timeout)
            if proc.returncode == 0 and dest.is_file():
                return 'powerpoint'
            errors.append('PowerPoint: ' + (proc.stderr.decode('utf-8', errors='replace')[-1800:] or 'No PDF output'))
        except subprocess.TimeoutExpired:
            errors.append('PowerPoint timed out; verify any remaining Office process before retrying')
    binary = find_soffice(soffice)
    if binary and engine in ('auto', 'libreoffice'):
        with tempfile.TemporaryDirectory(prefix='pbs-lo-') as profile:
            proc = subprocess.run([binary, '-env:UserInstallation=' + Path(profile).as_uri(), '--headless',
                                   '--convert-to', 'pdf', '--outdir', str(dest.parent), str(source)],
                                  capture_output=True, timeout=timeout)
        expected = dest.parent / (source.stem + '.pdf')
        if proc.returncode == 0 and expected.is_file():
            if expected != dest:
                shutil.copyfile(expected, dest)
            return 'libreoffice'
        errors.append('LibreOffice: ' + proc.stderr.decode('utf-8', errors='replace'))
    raise RuntimeError('No working PowerPoint/LibreOffice renderer. Install/configure one or supply a PDF genuinely exported from Office. ' + '\n'.join(errors))


def rasterize(pdf, output, dpi=144, columns=3):
    paths = []
    with pdfium.PdfDocument(str(pdf)) as doc:
        for i in range(len(doc)):
            page = doc[i]
            bitmap = page.render(scale=dpi / 72)
            im = bitmap.to_pil()
            path = output / f'slide-{i + 1:03}.png'
            im.save(path); paths.append(path)
            im.close(); bitmap.close(); page.close()
    if not paths:
        raise ValueError('Input has no pages')
    thumbs = []
    for path in paths:
        with Image.open(path) as original:
            im = original.convert('RGB'); im.thumbnail((480, 360))
            thumbs.append(im)
    cell_w = 500
    cell_h = max(im.height for im in thumbs) + 40
    columns = min(columns, len(paths))
    canvas = Image.new('RGB', (cell_w * columns, cell_h * math.ceil(len(paths) / columns)), '#E5E7EB')
    draw = ImageDraw.Draw(canvas)
    for i, im in enumerate(thumbs):
        x, y = (i % columns) * cell_w + 10, (i // columns) * cell_h + 25
        canvas.paste(im, (x, y)); draw.text((x, y - 18), f'{i + 1:03}', fill='#111111'); im.close()
    montage = output / 'contact-sheet.png'; canvas.save(montage); canvas.close()
    return paths, montage


def render(source, output, engine='auto', soffice=None, dpi=144, timeout=120, columns=3):
    source, output = Path(source).resolve(), Path(output).resolve()
    if not 36 <= dpi <= 600 or not 1 <= columns <= 8:
        raise ValueError('DPI must be 36..600; columns 1..8')
    output.mkdir(parents=True, exist_ok=True)
    if list(output.glob('slide-*.png')) or (output / 'render-report.json').exists():
        raise FileExistsError('Use a fresh preview directory to avoid stale/mixed pages')
    try:
        if source.suffix.lower() == '.pdf':
            pdf, used = source, 'pdf-input'
        else:
            pdf = output / 'office-render.pdf'
            used = office_pdf(source, pdf, engine, soffice, timeout)
        paths, montage = rasterize(pdf, output, dpi, columns)
        if source.suffix.lower() == '.pptx':
            from pptx import Presentation
            if len(paths) != len(Presentation(source).slides):
                raise RuntimeError('Renderer page count differs from source (check hidden slides)')
        report = {'status': 'rendered', 'source': str(source), 'engine': used, 'dpi': dpi,
                  'pages': [str(p) for p in paths], 'contact_sheet': str(montage),
                  'visual_review': 'required; rendering alone is not approval'}
    except Exception as exc:
        report = {'status': 'blocked', 'source': str(source), 'error': str(exc)}
        (output / 'render-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        raise
    (output / 'render-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


def main():
    utf8_stdio()
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('source', nargs='?', type=Path)
    p.add_argument('-o', '--output', type=Path)
    p.add_argument('--engine', choices=['auto', 'powerpoint', 'libreoffice'], default='auto')
    p.add_argument('--soffice')
    p.add_argument('--dpi', type=int, default=144)
    p.add_argument('--timeout', type=int, default=120)
    p.add_argument('--columns', type=int, default=3)
    p.add_argument('--probe', action='store_true')
    a = p.parse_args()
    if a.probe:
        print(json.dumps({'libreoffice': find_soffice(a.soffice), 'powerpoint': 'COM runtime check required' if os.name == 'nt' else 'unsupported'}, ensure_ascii=False))
        return
    if not a.source or not a.output:
        p.error('source and --output are required unless --probe')
    try:
        r = render(a.source, a.output, a.engine, a.soffice, a.dpi, a.timeout, a.columns)
        print(json.dumps({'status':r['status'], 'engine':r['engine'], 'page_count':len(r['pages']),
                          'report':str(a.output.resolve() / 'render-report.json'),
                          'review':'Inspect actual renders at readable size; review_proxy is optional. IF API Desktop requires final-body preflight, not a fixed thumbnail cap.'}, ensure_ascii=False))
    except Exception as exc:
        print(str(exc)); raise SystemExit(2)

if __name__ == '__main__':
    main()
