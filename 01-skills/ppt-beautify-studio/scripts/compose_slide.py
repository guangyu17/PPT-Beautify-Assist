#!/usr/bin/env python3
"""Build native editable slides from layout JSON. Coordinates: in/pt/px; font sizes always pt.
Usage: compose_slide.py layout.json -o new.pptx [--replace-generated]
"""
import argparse
import json
import os
import re
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement
from layout_components import expand_spec
from _common import (Units, hex_to_rgb, normalize_paragraphs, estimate_text_block,
                     cover_crop, contain_box, resolve_path, utf8_stdio)


def rgb(value):
    return RGBColor(*hex_to_rgb(value))


def fill(obj, color, alpha=1):
    if not 0 <= alpha <= 1:
        raise ValueError('Fill alpha must be 0..1')
    if color is None:
        obj.fill.background()
        return
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(color)
    solid = obj.fill._xPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
    if solid is not None and len(solid) and alpha < 1:
        node = OxmlElement('a:alpha')
        node.set('val', str(round(max(0, alpha) * 100000)))
        solid[0].append(node)



def line_style(shape, style):
    if not style:
        shape.line.fill.background()
        return
    width = float(style.get('width_pt', .8))
    alpha = float(style.get('alpha', 1))
    if not 0 < width <= 20 or not 0 <= alpha <= 1:
        raise ValueError('Line width must be 0..20pt, alpha 0..1')
    shape.line.color.rgb = rgb(style.get('color', '#254453'))
    shape.line.width = Pt(width)
    ln = shape.line._get_or_add_ln()
    if alpha < 1:
        color = ln.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        if color is not None and len(color):
            a = OxmlElement('a:alpha'); a.set('val', str(round(alpha * 100000))); color[0].append(a)
    dash = style.get('dash', 'solid')
    if dash not in ('solid','dash','dot','dashDot','longDash'):
        raise ValueError('Unsupported line dash')
    for old in list(ln):
        if old.tag.endswith('}prstDash'):ln.remove(old)
    node = OxmlElement('a:prstDash');node.set('val',dash);ln.append(node)


def native_shadow(shape):
    # Same restrained shadow vocabulary as the existing picture path; no rasterization.
    sppr = shape._element.spPr
    existing = sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
    if existing is not None:sppr.remove(existing)
    effects=OxmlElement('a:effectLst');shadow=OxmlElement('a:outerShdw')
    for k,v in {'blurRad':'63500','dist':'25400','dir':'5400000','algn':'ctr','rotWithShape':'0'}.items():shadow.set(k,v)
    color=OxmlElement('a:srgbClr');color.set('val','000000')
    alpha=OxmlElement('a:alpha');alpha.set('val','14000')
    color.append(alpha);shadow.append(color);effects.append(shadow);sppr.append(effects)


def add_connector(slide, el):
    route={'straight':MSO_CONNECTOR.STRAIGHT,'elbow':MSO_CONNECTOR.ELBOW}[el.get('route','straight')]
    shape=slide.shapes.add_connector(route,*[Inches(el[k]) for k in ('x1','y1','x2','y2')])
    line_style(shape,el.get('line',{'color':'#254453','width_pt':.8}))
    ln=shape.line._get_or_add_ln()
    for attribute,tag in (('start_arrow','a:headEnd'),('end_arrow','a:tailEnd')):
        value=el.get(attribute,'none')
        if value not in ('none','triangle','stealth','diamond','oval','arrow'):
            raise ValueError('Unsupported arrowhead: '+value)
        node=OxmlElement(tag);node.set('type',value);node.set('w','sm');node.set('len','sm');ln.append(node)
    return shape


def font(run, style):
    f = run.font
    f.size = Pt(float(style.get('size') or 18))
    f.name = style.get('latin_font') or style.get('font') or 'Arial'
    f.bold = bool(style.get('bold'))
    f.italic = bool(style.get('italic'))
    f.underline = bool(style.get('underline'))
    f.color.rgb = rgb(style.get('color') or '#1F2937')
    props = run._r.get_or_add_rPr()
    east = OxmlElement('a:ea')
    east.set('typeface', style.get('font') or 'Microsoft YaHei')
    props.append(east)


def add_text(shape, el, defaults):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = bool(el.get('wrap', True))
    tf.auto_size = MSO_AUTO_SIZE.NONE
    margin = float(el.get('margin', 0.05))
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}[el.get('valign', 'top')]
    pars = normalize_paragraphs(el, defaults)
    if el.get('autofit') == 'shrink':
        w, h = shape.width / 12700 - margin * 144, shape.height / 12700 - margin * 144
        floor = float(el.get('min_size', 10))
        while estimate_text_block(pars, w, tf.word_wrap)['height_pt'] > h:
            sizes = [r['size'] for p in pars for r in p['runs']]
            if not sizes or min(sizes) <= floor:
                break
            for p in pars:
                for r in p['runs']:
                    r['size'] = max(floor, r['size'] - 0.5)
    for i, pd in enumerate(pars):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}[pd.get('align') or 'left']
        p.level = min(8, max(0, pd.get('level', 0)))
        p.line_spacing = float(pd.get('line_spacing') or 1.0)
        p.space_before = Pt(float(pd.get('space_before') or 0))
        p.space_after = Pt(float(pd.get('space_after') or 0))
        props = p._p.get_or_add_pPr()
        indent = p.level * 27 + (18 if pd.get('bullet') else 0)
        props.set('marL', str(round(indent * 12700)))
        props.set('indent', str(-round(14 * 12700)) if pd.get('bullet') else '0')
        bullet = OxmlElement('a:buChar' if pd.get('bullet') else 'a:buNone')
        if pd.get('bullet'):
            bullet.set('char', '•')
        props.append(bullet)
        for rd in pd['runs']:
            for index, segment in enumerate(re.split(r'[\n\v]', rd['text'])):
                if index:
                    p.add_line_break()
                r = p.add_run()
                r.text = segment
                font(r, rd)
    if el.get('fill'):
        fill(shape, el['fill'], float(el.get('fill_alpha', 1)))


def picture(slide, path, box, fit='cover', el=None):
    el = el or {}
    with Image.open(path) as image:
        iw, ih = image.size
    x, y, w, h = box
    if fit == 'stretch':
        raise ValueError('Image stretching is forbidden; use cover or contain')
    if fit == 'contain':
        x, y, w, h = contain_box(iw, ih, x, y, w, h)
    elif fit != 'cover':
        raise ValueError('Unknown image fit: ' + fit)
    pic = slide.shapes.add_picture(str(path), *map(Inches, (x, y, w, h)))
    if fit == 'cover':
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cover_crop(iw, ih, w, h)
    if el.get('mask_shape'):
        if el['mask_shape'] not in ('ellipse','rounded_rect'):
            raise ValueError('mask_shape: ellipse or rounded_rect')
        pic._element.spPr.prstGeom.set('prst', 'ellipse' if el['mask_shape']=='ellipse' else 'roundRect')
    if el.get('radius'):
        geom = pic._element.spPr.prstGeom
        geom.set('prst', 'roundRect')
        av = geom.find('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
        guide = OxmlElement('a:gd')
        guide.set('name', 'adj')
        guide.set('fmla', f'val {round(min(0.5, float(el["radius"]) / min(w, h)) * 100000)}')
        av.append(guide)
    if el.get('border'):
        pic.line.color.rgb = rgb(el['border']['color'])
        pic.line.width = Pt(el['border'].get('width_pt', 1))
    if el.get('shadow'):
        effects = OxmlElement('a:effectLst')
        shadow = OxmlElement('a:outerShdw')
        for k, v in {'blurRad': '63500', 'dist': '25400', 'dir': '5400000', 'algn': 'ctr', 'rotWithShape': '0'}.items():
            shadow.set(k, v)
        color = OxmlElement('a:srgbClr'); color.set('val', '000000')
        alpha = OxmlElement('a:alpha'); alpha.set('val', '18000')
        color.append(alpha); shadow.append(color); effects.append(shadow)
        pic._element.spPr.append(effects)
    return pic


def compose(spec_path, output, replace=False):
    spec_path, output = Path(spec_path).resolve(), Path(output).resolve()
    raw_spec = json.loads(spec_path.read_text(encoding='utf-8-sig'))
    if raw_spec.get('studio_profile'):
        for slide_data in raw_spec.get('slides', []):
            for element in slide_data.get('elements', []):
                if element.get('type')=='panel' and element.get('texture'):
                    texture = resolve_path(element['texture'], spec_path.parent)
                    if not texture.is_file():
                        raise FileNotFoundError('Generated panel texture missing; do not silently downgrade to native panel')
    spec = expand_spec(raw_spec)
    if output.exists() and not replace:
        raise FileExistsError(f'Refusing to overwrite {output}; choose a new output')
    if output == spec_path or (spec.get('source') and output == resolve_path(spec['source'], spec_path.parent)):
        raise ValueError('Cannot overwrite source')
    if output.exists() and replace:
        if Presentation(output).core_properties.identifier != 'ppt-beautify-studio-generated-v1':
            raise ValueError('Replacement refused: output is not marked as generated by this skill')
    units = Units(spec)
    defaults = {'font': 'Microsoft YaHei', 'latin_font': 'Arial', 'size': 18, 'color': '#1F2937', **spec.get('defaults', {})}
    prs = Presentation()
    prs.core_properties.identifier = 'ppt-beautify-studio-generated-v1'
    prs.slide_width, prs.slide_height = Inches(units.sw), Inches(units.sh)
    if not spec.get('slides'):
        raise ValueError('Layout must have at least one slide')
    for sn, sd in enumerate(spec['slides'], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fill(slide.background, sd.get('background_color', '#FFFFFF'))
        if sd.get('background'):
            pic = picture(slide, resolve_path(sd['background'], spec_path.parent), (0, 0, units.sw, units.sh), sd.get('background_fit', 'cover'))
            pic.name = 'background'
        for en, el in enumerate(sd.get('elements', []), 1):
            typ = el['type']
            box = units.box(el)
            if typ != 'connector' and (box[2] <= 0 or box[3] <= 0):
                raise ValueError(f'Slide {sn} element {en}: width and height must be positive')
            pos = list(map(Inches, box))
            if typ == 'connector':
                shape = add_connector(slide, el)
            elif typ == 'text':
                shape = slide.shapes.add_textbox(*pos)
                add_text(shape, el, defaults)
            elif typ == 'image':
                shape = picture(slide, resolve_path(el['path'], spec_path.parent), box, el.get('fit', 'cover'), el)
            elif typ == 'shape':
                kind = {'rect': MSO_SHAPE.RECTANGLE, 'rounded_rect': MSO_SHAPE.ROUNDED_RECTANGLE, 'ellipse': MSO_SHAPE.OVAL, 'diamond': MSO_SHAPE.DIAMOND, 'star4': MSO_SHAPE.STAR_4_POINT, 'star8': MSO_SHAPE.STAR_8_POINT, 'snip_round_rect': MSO_SHAPE.SNIP_ROUND_RECTANGLE, 'snip_diag_rect': MSO_SHAPE.SNIP_2_DIAG_RECTANGLE, 'round_diag_rect': MSO_SHAPE.ROUND_2_DIAG_RECTANGLE}.get(el.get('shape', 'rect'))
                if kind is None:
                    raise ValueError('Unsupported native shape')
                shape = slide.shapes.add_shape(kind, *pos)
                fill(shape, el.get('fill'), float(el.get('alpha', 1)))
                line_style(shape, el.get('line'))
                if el.get('shadow'):
                    native_shadow(shape)
            elif typ == 'table':
                rows = el['rows']
                if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
                    raise ValueError('Table rows must be a nonempty rectangular array')
                shape = slide.shapes.add_table(len(rows), len(rows[0]), *pos)
                for ri, row in enumerate(rows):
                    for ci, text in enumerate(row):
                        cell = shape.table.cell(ri, ci)
                        header = ri == 0 and el.get('header', True)
                        fill(cell, el.get('header_fill', '#1E3A5F') if header else el.get('fill', '#FFFFFF'))
                        te = {'text': str(text), 'size': el.get('size', 14), 'bold': header,
                              'color': el.get('header_color', '#FFFFFF') if header else el.get('color', defaults['color'])}
                        add_text(cell, te, defaults)
                for merge in el.get('merges', []):
                    r1, c1, r2, c2 = merge
                    shape.table.cell(r1, c1).merge(shape.table.cell(r2, c2))
            elif typ == 'chart':
                kinds = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
                         'line': XL_CHART_TYPE.LINE, 'pie': XL_CHART_TYPE.PIE, 'doughnut': XL_CHART_TYPE.DOUGHNUT,
                         'area': XL_CHART_TYPE.AREA, 'stacked_column': XL_CHART_TYPE.COLUMN_STACKED,
                         'stacked_bar': XL_CHART_TYPE.BAR_STACKED}
                data = CategoryChartData(); data.categories = el['categories']
                for series in el['series']:
                    if len(series['values']) != len(el['categories']):
                        raise ValueError('Chart category/value lengths differ')
                    data.add_series(series['name'], series['values'])
                shape = slide.shapes.add_chart(kinds[el.get('chart_type', 'column')], *pos, data)
                chart = shape.chart
                chart.font.size = Pt(el.get('size', 12))
                chart.font.name = defaults['font']
                chart.has_legend = bool(el.get('legend', True))
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                for i, series in enumerate(chart.series):
                    colors = el.get('colors') or []
                    if colors:
                        series.format.fill.solid(); series.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
                        series.format.line.color.rgb = rgb(colors[i % len(colors)])
            else:
                raise ValueError(f'Unsupported element type: {typ}; do not silently flatten it')
            # Stock Office shape styles can add a theme shadow even when none was
            # requested; on translucent panels it can turn a pale glass surface gray.
            for ref in shape._element.xpath('.//a:effectRef'):
                ref.set('idx', '0')
            if not el.get('shadow') and typ in ('shape','connector','image'):
                sppr = shape._element.spPr
                effects = sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
                if effects is not None:
                    sppr.remove(effects)
                sppr.append(OxmlElement('a:effectLst'))
            if typ == 'shape' and el.get('shape') in ('star4','star8') and len(shape.adjustments):
                shape.adjustments[0] = .16
            shape.name = str(el.get('id') or f's{sn:03}-e{en:03}-{typ}')
            if el.get('rotation'):
                shape.rotation = float(el['rotation'])
        if sd.get('notes'):
            slide.notes_slide.notes_text_frame.text = sd['notes']
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(suffix='.pptx', dir=output.parent); os.close(fd)
    try:
        prs.save(temp)
        os.replace(temp, output)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
    return {'slides': len(prs.slides), 'output': str(output), 'editable': ['text', 'pictures', 'tables', 'charts']}


def main():
    utf8_stdio()
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('layout', type=Path)
    p.add_argument('-o', '--output', type=Path, required=True)
    p.add_argument('--replace-generated', action='store_true', help='Only use for a previously generated output, never an original')
    a = p.parse_args()
    print(json.dumps(compose(a.layout, a.output, a.replace_generated), ensure_ascii=False))

if __name__ == '__main__':
    main()
