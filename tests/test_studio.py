"""Portable offline regression suite; no services, user settings, installs or real credentials."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
SCRIPTS=ROOT/'skills/ppt-beautify-studio/scripts'
sys.path.insert(0,str(SCRIPTS))
from studio_settings import load, update
from request_preflight import inspect_request
from layout_components import expand_spec
from compose_slide import compose
from check_layout import lint
from blend_image import gradient_field, blend_recipe
from review_proxy import prepare
from PIL import Image
from pptx import Presentation


def write_json(path,data):
    path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    return path


class StudioTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.prefs=self.root/'preferences.json'
    def tearDown(self):self.tmp.cleanup()

    def test_native_default_without_creating_preferences(self):
        self.assertEqual(load(self.prefs)['backend'],'native')
        self.assertIsNone(load(self.prefs)['tier'])
        self.assertFalse(self.prefs.exists())
        r=inspect_request(preferences=self.prefs)
        self.assertEqual(r['status'],'exempt')
        self.assertEqual(r['network_requests_sent'],0)

    def test_persistent_choice_and_explicit_consent(self):
        with self.assertRaises(ValueError):update(self.prefs,backend='relay')
        cfg=self.root/'private.json';cfg.write_text('{}')
        update(self.prefs,backend='relay',tier='high',relay_config=cfg,confirmed=True)
        self.assertEqual(load(self.prefs)['tier'],'high')
        self.assertEqual(load(self.prefs)['relay_config_path'],str(cfg.resolve()))
        update(self.prefs,backend='native',confirmed=True)
        self.assertEqual(load(self.prefs)['tier'],'high')
        with self.assertRaises(ValueError):update(self.prefs,backend='invalid',confirmed=True)
        self.assertEqual(load(self.prefs)['backend'],'native')

    def test_cannot_turn_off_guard_by_version(self):
        with self.assertRaises(ValueError):update(self.prefs,guard='off',confirmed=True)
        update(self.prefs,guard='off',evidence='Test fixture: explicit fix confirmation',confirmed=True)
        self.assertEqual(inspect_request('ifapi',preferences=self.prefs)['status'],'exempt')
        update(self.prefs,guard='on',confirmed=True)
        self.assertEqual(inspect_request('ifapi',preferences=self.prefs)['status'],'needs_capture')

    def test_exact_body_boundary_and_hash(self):
        update(self.prefs,budget=1000,confirmed=True)
        body=self.root/'final-body.bin';body.write_bytes(b'x'*1000)
        r=inspect_request('ifapi',body=body,complete=True,preferences=self.prefs)
        self.assertEqual(r['status'],'pass');self.assertEqual(r['measured_bytes'],1000)
        self.assertEqual(r['body_sha256'],hashlib.sha256(body.read_bytes()).hexdigest())
        body.write_bytes(b'x'*1001)
        self.assertEqual(inspect_request('ifapi',body=body,complete=True,preferences=self.prefs)['status'],'oversize')
        self.assertEqual(body.stat().st_size,1001)

    def test_63mb_rejection_fixture_streaming(self):
        # Sparse-size fixture is read in bounded chunks; never printed or transmitted.
        body=self.root/'large-body.bin'
        with body.open('wb') as f:f.truncate(63000000)
        r=inspect_request('ifapi',body=body,complete=True,preferences=self.prefs)
        self.assertEqual(r['status'],'oversize');self.assertEqual(r['measured_bytes'],63000000)
        self.assertEqual(r['budget_bytes'],60000000)

    def test_partial_and_estimate_never_pass_ifapi(self):
        body=write_json(self.root/'args.json',{'prompt':'small prompt','reference_path':'local.png'})
        self.assertEqual(inspect_request('ifapi',body=body,preferences=self.prefs)['status'],'needs_capture')
        asset=self.root/'input.png';asset.write_bytes(b'abcdefgh')
        manifest=write_json(self.root/'estimate.json',{'unknown_history':False,'text_bytes':100,
                 'assets':[{'path':'input.png','encoding':'base64','occurrences':2}]})
        r=inspect_request('ifapi',manifest=manifest,preferences=self.prefs)
        self.assertEqual(r['estimated_known_bytes'],124)
        self.assertEqual(r['status'],'needs_capture')
        update(self.prefs,budget=120,confirmed=True)
        self.assertEqual(inspect_request('ifapi',manifest=manifest,preferences=self.prefs)['status'],'oversize')
        with self.assertRaises(ValueError):inspect_request('ifapi',manifest=manifest,complete=True,preferences=self.prefs)

    def test_relay_transport_explicit_and_native_exempt(self):
        self.assertEqual(inspect_request('relay',preferences=self.prefs)['status'],'exempt')
        self.assertEqual(inspect_request('relay',transport='ifapi-desktop',preferences=self.prefs)['status'],'needs_capture')
        body=self.root/'small.bin';body.write_bytes(b'xx')
        update(self.prefs,budget=1,confirmed=True)
        self.assertEqual(inspect_request('native',body=body,complete=True,preferences=self.prefs)['status'],'exempt')

    def test_preferences_reject_secrets(self):
        d=load(self.prefs);d['api_key']='not-a-real-credential'
        write_json(self.prefs,d)
        with self.assertRaises(ValueError):load(self.prefs)

    def panel_spec(self,tier='high'):
        return {'studio_profile':{'tier':tier,'backend':'native'},'aspect':'16:9',
                'slides':[{'elements':[{'type':'panel','x':1,'y':1,'w':6,'h':2,'material':'acrylic',
                'text':'Editable text','size':22,'alpha':.06,'padding':.2}]}]}

    def test_generated_panel_gate_all_tiers(self):
        for tier in ('ultra','high','medium','low'):
            with self.subTest(tier=tier),self.assertRaises(ValueError):expand_spec(self.panel_spec(tier))
        s=self.panel_spec();s['slides'][0]['elements'][0].update(material_source='generated',texture='missing.png')
        with self.assertRaises(FileNotFoundError):compose(write_json(self.root/'missing.json',s),self.root/'missing.pptx')
        self.assertFalse((self.root/'missing.pptx').exists())

    def test_panel_native_text_and_legacy_compatibility(self):
        # Synthetic fixture: this tests the schema, NOT provenance of AI generation.
        image=self.root/'fixture.png';Image.new('RGBA',(600,200),(220,230,240,100)).save(image)
        s=self.panel_spec();e=s['slides'][0]['elements'][0]
        e.pop('alpha')
        e.update(texture=str(image),material_source='generated',inset=0)
        expanded=expand_spec(s)
        surface=next(v for v in expanded['slides'][0]['elements'] if v['id'].endswith('-surface'))
        self.assertEqual(surface['alpha'],0)
        self.assertIsNone(surface['line'])
        path=write_json(self.root/'layout.json',s);out=self.root/'output.pptx';compose(path,out)
        p=Presentation(out)
        self.assertTrue(any(sh.has_text_frame and sh.text=='Editable text' for sh in p.slides[0].shapes))
        self.assertEqual(sum(sh.shape_type==13 for sh in p.slides[0].shapes),1)
        with self.assertRaises(FileExistsError):compose(path,out)
        legacy=self.panel_spec();legacy.pop('studio_profile')
        self.assertTrue(expand_spec(legacy)['slides'][0]['elements'])

    def test_readable_proxies_no_old_ceiling(self):
        source=self.root/'source.png';Image.new('RGB',(2200,1400),'#DDEEFF').save(source)
        r=prepare(source,self.root/'view.jpg',max_edge=1800,max_kib=600)
        self.assertEqual(max(r['pixels']),1800)
        self.assertLessEqual(r['bytes'],600*1024)

    def test_gradient_and_native_connector_regression(self):
        a=gradient_field(100,60,{'angle':90,'stops':[[0,0],[1,1]]})
        self.assertAlmostEqual(a[0,0],0);self.assertAlmostEqual(a[-1,0],1)
        s={'slides':[{'elements':[{'type':'connector','id':'line','x1':1,'y1':1,'x2':5,'y2':1,
             'end_arrow':'arrow'},{'type':'text','text':'Not crossed','x':1,'y':2,'w':5,'h':1,'size':20}]}]}
        path=write_json(self.root/'line.json',s);compose(path,self.root/'line.pptx')
        self.assertEqual(lint(path)['status'],'no_detected_errors')


if __name__=='__main__':unittest.main()
