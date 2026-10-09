#!/usr/bin/env python3
"""Read PPTX content and original media; enrich with ppt-master's read-only intake.
Usage: extract_pptx.py source.pptx -o extracted/deck.json [--media DIR]
"""
import argparse
import hashlib
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from _common import utf8_stdio


def color_info(color):
    try:
        if color.type is None:
            return None
        if color.type == 1:
            return {'rgb': '#' + str(color.rgb), 'brightness': color.brightness}
        return {'type': str(color.type), 'theme': str(color.theme_color), 'brightness': color.brightness}
    except (AttributeError, ValueError):
        return {'unresolved': True}


def font_info(font):
    return {'font': font.name, 'size_pt': font.size.pt if font.size else None,
            'bold': font.bold, 'italic': font.italic, 'underline': font.underline,
            'color': color_info(font.color),
            'script_fonts': {node.tag.rsplit('}', 1)[-1]: node.get('typeface')
                             for node in font._rPr if node.tag.rsplit('}', 1)[-1] in ('latin', 'ea', 'cs')}}


def paragraphs(tf):
    result = []
    for p in tf.paragraphs:
        result.append({'text': p.text, 'level': p.level,
                       'alignment': str(p.alignment) if p.alignment else None,
                       'default_style': font_info(p.font),
                       'runs': [{'text': r.text, **font_info(r.font)} for r in p.runs],
                       'xml': p._p.xml})
    return result


def extract(source, output, media, master):
    source, output, media = Path(source).resolve(), Path(output).resolve(), Path(media).resolve()
    if source == output:
        raise ValueError('Output must not overwrite source')
    output.parent.mkdir(parents=True, exist_ok=True)
    media.mkdir(parents=True, exist_ok=True)
    prs = Presentation(source)
    result = {'schema_version': 1, 'source': str(source),
              'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
              'slide_size': {'width': prs.slide_width / 914400, 'height': prs.slide_height / 914400},
              'units': 'in', 'slides': [], 'warnings': [],
              'style_contract': 'Null means inherited, not a default. Consult intake, XML and rendered page before reauthoring.'}
    # Use the installed resolver/intake instead of duplicating its theme/chart/SmartArt machinery.
    master = Path(master).resolve()
    for script, args in [('attribution_guard.py', []),
                         ('pptx_intake.py', [str(source), '-o', str(output.parent / 'intake')])]:
        proc = subprocess.run([sys.executable, '-B', str(master / 'scripts' / script), *args],
                              capture_output=True, text=True, encoding='utf-8', errors='replace')
        if proc.returncode:
            raise RuntimeError(f'ppt-master {script} failed; stopped without bypassing guard: {proc.stderr or proc.stdout}')
    identity = output.parent / 'intake' / (source.stem + '.identity.json')
    result['identity'] = json.loads(identity.read_text(encoding='utf-8'))
    result['intake'] = str(output.parent / 'intake')

    def walk(shapes, page, transform=(1, 1, 0, 0), parent=None):
        sx, sy, tx, ty = transform
        items = []
        for z, shape in enumerate(shapes):
            key = f'{parent}_{shape.shape_id}' if parent else str(shape.shape_id)
            item = {'id': key, 'name': shape.name, 'z': z, 'type': str(shape.shape_type),
                    'x': (shape.left * sx + tx) / 914400, 'y': (shape.top * sy + ty) / 914400,
                    'w': shape.width * sx / 914400, 'h': shape.height * sy / 914400,
                    'rotation': shape.rotation, 'xml': shape._element.xml}
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                xf = shape._element.grpSpPr.xfrm
                if xf is not None and xf.chExt is not None and xf.chExt.cx and xf.chExt.cy:
                    xs, ys = shape.width / xf.chExt.cx, shape.height / xf.chExt.cy
                    nt = (sx * xs, sy * ys, tx + sx * (shape.left - xf.chOff.x * xs), ty + sy * (shape.top - xf.chOff.y * ys))
                else:
                    nt = transform
                item['children'] = walk(shape.shapes, page, nt, key)
                if shape.rotation or 'flipH="1"' in shape._element.xml or 'flipV="1"' in shape._element.xml:
                    result['warnings'].append(f'Page {page} group {key}: rotated/flipped group needs visual coordinate verification')
            if shape.has_text_frame:
                item['text'] = shape.text_frame.text
                item['paragraphs'] = paragraphs(shape.text_frame)
                item['margins'] = {k: getattr(shape.text_frame, 'margin_' + k) / 914400 for k in ('left', 'right', 'top', 'bottom')}
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                path = media / f's{page:03}_{key}.{image.ext}'
                path.write_bytes(image.blob)
                item.update(image=str(path), pixel_size=list(image.size), dpi=list(image.dpi),
                            crop={k: getattr(shape, 'crop_' + k) for k in ('left', 'top', 'right', 'bottom')})
            if shape.has_table:
                item['table'] = [[{'text': c.text, 'paragraphs': paragraphs(c.text_frame),
                                   'is_merge_origin': c.is_merge_origin, 'is_spanned': c.is_spanned,
                                   'span_width': c.span_width, 'span_height': c.span_height}
                                  for c in row.cells] for row in shape.table.rows]
            if shape.has_chart:
                chart = shape.chart
                item['chart'] = {'chart_type': str(chart.chart_type), 'xml': chart._chartSpace.xml, 'series': []}
                for series in chart.series:
                    try:
                        item['chart']['series'].append({'name': series.name, 'values': list(series.values)})
                    except (AttributeError, ValueError) as exc:
                        result['warnings'].append(f'Page {page} chart {key}: {exc}; use intake/XML')
                try:
                    item['chart']['categories'] = [c.label for c in chart.plots[0].categories]
                except (AttributeError, ValueError):
                    item['chart']['categories'] = None
            items.append(item)
        return items

    for n, slide in enumerate(prs.slides, 1):
        result['slides'].append({'index': n, 'hidden': slide._element.get('show') == '0',
                                 'elements': walk(slide.shapes, n),
                                 'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None})
    # Preserve source XML evidence, including inherited master/layout text and embedded chart workbooks.
    with ZipFile(source) as archive:
        evidence = output.parent / 'ooxml'
        for name in archive.namelist():
            if name.startswith(('ppt/theme/', 'ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/embeddings/', 'ppt/media/')) and not name.endswith('/'):
                dest = (evidence / name).resolve()
                if not dest.is_relative_to(evidence.resolve()):
                    raise ValueError('Unsafe archive member')
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(archive.read(name))
        result['ooxml_evidence'] = str(evidence)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    return result


def main():
    utf8_stdio()
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('source', type=Path)
    p.add_argument('-o', '--output', required=True, type=Path)
    p.add_argument('--media', type=Path)
    p.add_argument('--ppt-master', type=Path, default=Path(__file__).resolve().parents[2] / 'ppt-master')
    a = p.parse_args()
    r = extract(a.source, a.output, a.media or a.output.parent / 'media', a.ppt_master)
    print(json.dumps({'slides': len(r['slides']), 'warnings': r['warnings'], 'output': str(a.output)}, ensure_ascii=False))

if __name__ == '__main__':
    main()
