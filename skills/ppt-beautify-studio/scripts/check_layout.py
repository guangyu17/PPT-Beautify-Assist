#!/usr/bin/env python3
"""Conservative layout lint, not a substitute for inspecting Office-rendered pages.
Exit 1: detected errors; 0: no detected errors (warnings/visual checks may remain).
"""
import argparse
from collections import Counter
import json
from pathlib import Path
import numpy as np
from PIL import Image, ImageOps
from layout_components import expand_spec, segment_intersects_box
from _common import (Units, normalize_paragraphs, estimate_text_block, overlap_area,
                     resolve_path, hex_to_rgb, rel_luminance, contrast_from_lum, utf8_stdio)


def content(prs):
    def shapes_text(shapes):
        strings, numbers = [], []
        for s in shapes:
            if s.shape_type == 6:
                ss, nn = shapes_text(s.shapes); strings += ss; numbers += nn
            if s.has_text_frame:
                strings += [p.text for p in s.text_frame.paragraphs if p.text]
            if s.has_table:
                strings += [c.text for row in s.table.rows for c in row.cells if c.text]
            if s.has_chart:
                for series in s.chart.series:
                    strings.append(series.name)
                    numbers += [str(v) for v in series.values]
                for plot in s.chart.plots:
                    try:
                        strings += [str(c.label) for c in plot.categories]
                    except AttributeError:
                        pass
        return strings, numbers
    return [(Counter(t), Counter(n)) for t, n in (shapes_text(slide.shapes) for slide in prs.slides)]


def background(slide, units, base, override=None):
    size = (1200, round(1200 * units.sh / units.sw))
    canvas = Image.new('RGB', size, slide.get('background_color', '#FFFFFF'))
    if slide.get('background'):
        path = resolve_path(slide['background'], base)
        if not path.exists() and override:
            path = Path(override) / Path(slide['background']).name
        with Image.open(path) as img:
            rgba = img.convert('RGBA')
            if slide.get('background_fit', 'cover') == 'cover':
                thumb = ImageOps.fit(rgba, size)
            else:
                thumb = ImageOps.contain(rgba, size)
            position = ((size[0] - thumb.width) // 2, (size[1] - thumb.height) // 2)
            # A feathered PNG exposes the native slide color. Dropping its alpha would
            # incorrectly assess hidden RGB pixels as the visible text background.
            canvas.paste(thumb, position, thumb.getchannel('A'))
    return canvas


def lint(layout, pptx=None, source_pptx=None, backgrounds=None):
    layout = Path(layout).resolve()
    spec = expand_spec(json.loads(layout.read_text(encoding='utf-8-sig')))
    units = Units(spec)
    defaults = {'font': 'Microsoft YaHei', 'size': 18, 'color': '#1F2937', **spec.get('defaults', {})}
    issues = []
    def add(level, code, slide, element, detail):
        issues.append(dict(level=level, code=code, slide=slide, element=element, detail=detail))
    for sn, slide in enumerate(spec.get('slides', []), 1):
        try:
            bg = background(slide, units, layout.parent, backgrounds)
        except Exception as exc:
            bg = None; add('error', 'background_missing', sn, None, str(exc))
        elements = slide.get('elements', [])
        boxes = []
        for en, el in enumerate(elements, 1):
            box = units.box(el); x, y, w, h = box
            boxes.append(box)
            if (el['type'] != 'connector' and (w <= 0 or h <= 0)) or x < -0.005 or y < -0.005 or x + w > units.sw + 0.005 or y + h > units.sh + 0.005:
                add('error', 'out_of_bounds', sn, en, f'box={box}, canvas={units.sw, units.sh}')
            if el['type']=='connector' and el.get('route')=='elbow':
                add('warning','elbow_route_review',sn,en,'Elbow path and arrow clearance require actual rendered review')
            if el.get('rotation'):
                add('warning', 'rotated_bounds', sn, en, 'Rotated corners and text need visual review')
            if el['type'] == 'image':
                path = resolve_path(el['path'], layout.parent)
                try:
                    with Image.open(path) as im:
                        dpi = min(im.width / max(w, 0.001), im.height / max(h, 0.001))
                    if dpi < 150:
                        add('warning', 'image_resolution', sn, en, f'Conservative effective resolution {dpi:.0f} DPI')
                except Exception as exc:
                    add('error', 'image_missing', sn, en, str(exc))
                if el.get('fit') == 'stretch':
                    add('error', 'image_distortion', sn, en, 'Use cover/contain instead of stretch')
            if el['type'] not in ('text', 'table', 'chart', 'shape', 'image', 'connector'):
                add('error', 'unsupported_type', sn, en, el['type'])
            if el['type'] in ('table', 'chart'):
                if el.get('size', 14 if el['type'] == 'table' else 12) < 10:
                    add('error', 'minimum_font', sn, en, 'Table/chart font below 10pt')
                add('warning', 'native_data_visual_review', sn, en, 'Table cells/chart labels require native-rendered overflow and contrast review')
            if el['type'] != 'text':
                continue
            pars = normalize_paragraphs(el, defaults)
            margin = float(el.get('margin', 0.05)) * 72
            estimate = estimate_text_block(pars, max(1, w * 72 - 2 * margin), el.get('wrap', True))
            if estimate['height_pt'] > h * 72 - 2 * margin + 2 or estimate['max_line_width_pt'] > w * 72 - 2 * margin + 2:
                add('warning' if el.get('autofit') == 'shrink' else 'error', 'text_overflow_estimate', sn, en, f'{estimate}; available {w * 72 - 2 * margin:.1f} x {h * 72 - 2 * margin:.1f} pt (font estimate, verify actual render)')
            patch = None
            if bg is not None and w > 0 and h > 0:
                bx = (max(0, round(x / units.sw * bg.width)), max(0, round(y / units.sh * bg.height)),
                      min(bg.width, round((x + w) / units.sw * bg.width)), min(bg.height, round((y + h) / units.sh * bg.height)))
                if bx[2] > bx[0] and bx[3] > bx[1]:
                    patch = np.array(bg.crop(bx).resize((48, 24)), dtype=float)
                    if el.get('fill'):
                        alpha = float(el.get('fill_alpha', 1))
                        patch = patch * (1 - alpha) + np.array(hex_to_rgb(el['fill'])) * alpha
            prior_visual_layers = any(earlier['type'] in ('shape','image') and overlap_area(box, eb)>0
                                      for earlier, eb in zip(elements[:en-1], boxes[:-1]))
            own_opaque_fill = bool(el.get('fill')) and float(el.get('fill_alpha',1)) >= 1
            colors_done = set()
            for p in pars:
                for run in p['runs']:
                    size = float(run['size'])
                    if size < 10:
                        add('error', 'minimum_font', sn, en, f'{size}pt < 10pt')
                    elif size < 14 and el.get('role', 'body') == 'body':
                        add('warning', 'small_body_font', sn, en, f'{size}pt < recommended 14pt')
                    color = run.get('color') or '#1F2937'
                    threshold = 3 if size >= 18 or (size >= 14 and run.get('bold')) else 4.5
                    key = (color, threshold)
                    if patch is not None and key not in colors_done:
                        colors_done.add(key)
                        lum = np.array([rel_luminance(pixel) for pixel in patch.reshape(-1, 3)])
                        fg = rel_luminance(hex_to_rgb(color))
                        # The low percentile flags substantial low-contrast portions, not isolated antialias pixels.
                        ratios = np.array([contrast_from_lum(fg, b) for b in lum])
                        ratio = float(np.percentile(ratios, 10))
                        if ratio < threshold:
                            add('warning' if prior_visual_layers and not own_opaque_fill else 'error', 'low_contrast', sn, en, f'10th-percentile contrast {ratio:.2f}:1 < {threshold}:1 (background-only approximation; prior visual layers require rendered review)')
            if bg is not None:
                for earlier, eb in zip(elements[:en - 1], boxes[:-1]):
                    if overlap_area(box, eb) > 0:
                        add('warning', 'contrast_layers', sn, en, 'Earlier content overlaps text; background-only contrast cannot certify this text')
                        break
        for i, a in enumerate(elements):
            for j in range(i + 1, len(elements)):
                b = elements[j]
                aid, bid = a.get('id', str(i + 1)), b.get('id', str(j + 1))
                if bid in a.get('allow_overlap_with', []) or aid in b.get('allow_overlap_with', []):
                    continue
                if {a['type'], b['type']} == {'connector','text'}:
                    line, text_box = (a, boxes[j]) if a['type']=='connector' else (b, boxes[i])
                    if line.get('route','straight')=='straight' and segment_intersects_box(line['x1'],line['y1'],line['x2'],line['y2'],text_box):
                        add('error','connector_crosses_text',sn,f'{i+1},{j+1}','A straight native connector crosses a text box; route it around text or verify explicit intentional interaction')
                    continue
                area = overlap_area(boxes[i], boxes[j])
                if area < 0.01:
                    continue
                # Only explicit, pair-specific annotations waive an intentional overlap.
                aid, bid = a.get('id', str(i + 1)), b.get('id', str(j + 1))
                if bid in a.get('allow_overlap_with', []) or aid in b.get('allow_overlap_with', []):
                    continue
                level = 'error' if 'text' in (a['type'], b['type']) and a['type'] != 'shape' and b['type'] != 'shape' else 'warning'
                add(level, 'overlap', sn, f'{i + 1},{j + 1}', f'{area:.3f} sq.in; annotate intentional layering with allow_overlap_with IDs')
        if bg is not None:
            bg.close()
    if pptx:
        from pptx import Presentation
        prs = Presentation(pptx)
        if len(prs.slides) != len(spec['slides']):
            add('error', 'slide_count', None, None, 'Output and layout page counts differ')
        if source_pptx:
            original = content(Presentation(source_pptx)); actual = content(prs)
            if len(original) != len(actual):
                add('error', 'source_slide_count', None, None, 'Original/output page counts differ')
            for i, (old, new) in enumerate(zip(original, actual), 1):
                if old != new:
                    add('error', 'content_changed', i, None, f'Exact paragraph/cell/chart-label or value multisets differ. Missing text: {dict(old[0] - new[0])}; added: {dict(new[0] - old[0])}; check numbers too.')
    elif source_pptx:
        raise ValueError('--source-pptx requires --pptx')
    return {'status': 'errors' if any(i['level'] == 'error' for i in issues) else 'no_detected_errors',
            'issues': issues, 'required_manual_checks': ['Office rendered text fit, alignment and occlusion',
                 'Generated background contains no text/pseudo-text/logo', 'Original wording including inherited master/SmartArt text',
                 'Image identity, shape, crop and style consistency'],
            'limitations': 'Font-based overflow and background contrast are estimates, not certification. Inherited styles, table cells, charts and arbitrary rotated/layered objects require rendered review.'}


def main():
    utf8_stdio()
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('layout', type=Path)
    p.add_argument('-o', '--output', type=Path)
    p.add_argument('--backgrounds', type=Path)
    p.add_argument('--pptx', type=Path)
    p.add_argument('--source-pptx', type=Path)
    a = p.parse_args()
    r = lint(a.layout, a.pptx, a.source_pptx, a.backgrounds)
    report_path = a.output or a.layout.with_name(a.layout.stem + '.layout-report.json')
    if report_path.resolve() in (a.layout.resolve(), *(p.resolve() for p in (a.pptx, a.source_pptx) if p)):
        raise ValueError('Report cannot overwrite an input')
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding='utf-8')
    sample = [{**i, 'detail': str(i['detail'])[:250]} for i in r['issues'][:3]]
    print(json.dumps({'status':r['status'], 'issue_count':len(r['issues']),
                      'sample_issues':sample, 'report':str(report_path.resolve())}, ensure_ascii=False))
    raise SystemExit(1 if r['status'] == 'errors' else 0)

if __name__ == '__main__':
    main()
